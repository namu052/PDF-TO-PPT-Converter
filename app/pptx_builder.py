from __future__ import annotations

import io
import logging
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR_TYPE, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.color_utils import hex_to_rgb
from app.font_utils import resolve_font
from app.models import (
    ArrowType,
    ElementType,
    LineStyle,
    PresentationData,
    RenderStrategy,
    ShapeKind,
    SlideElement,
    TextParagraph,
)

logger = logging.getLogger(__name__)


SHAPE_KIND_TO_MSO = {
    ShapeKind.RECTANGLE: MSO_SHAPE.RECTANGLE,
    ShapeKind.ROUNDED_RECTANGLE: MSO_SHAPE.ROUNDED_RECTANGLE,
    ShapeKind.ELLIPSE: MSO_SHAPE.OVAL,
    ShapeKind.TRIANGLE: MSO_SHAPE.ISOSCELES_TRIANGLE,
    ShapeKind.RIGHT_ARROW: MSO_SHAPE.RIGHT_ARROW,
    ShapeKind.LEFT_ARROW: MSO_SHAPE.LEFT_ARROW,
    ShapeKind.UP_ARROW: MSO_SHAPE.UP_ARROW,
    ShapeKind.DOWN_ARROW: MSO_SHAPE.DOWN_ARROW,
    ShapeKind.DIAMOND: MSO_SHAPE.DIAMOND,
    ShapeKind.PENTAGON: MSO_SHAPE.PENTAGON,
    ShapeKind.HEXAGON: MSO_SHAPE.HEXAGON,
    ShapeKind.CLOUD: MSO_SHAPE.CLOUD,
    ShapeKind.STAR: MSO_SHAPE.STAR_5_POINT,
    ShapeKind.CHEVRON: MSO_SHAPE.CHEVRON,
    ShapeKind.CALLOUT: MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
}

ALIGN_MAP = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}


def pct_to_emu(pct: float, total_inches: float) -> int:
    return int(Inches((pct / 100.0) * total_inches))


def _bbox_to_emu(elem: SlideElement, w_in: float, h_in: float) -> tuple[int, int, int, int]:
    return (
        pct_to_emu(elem.bbox.x_pct, w_in),
        pct_to_emu(elem.bbox.y_pct, h_in),
        pct_to_emu(elem.bbox.width_pct, w_in),
        pct_to_emu(elem.bbox.height_pct, h_in),
    )


def _set_fill_opacity(shape, opacity: float) -> None:
    try:
        alpha = int(max(0.0, min(1.0, opacity)) * 100000)
        fill_elem = shape.fill._fill
        srgb = fill_elem.find(qn("a:srgbClr"))
        if srgb is None:
            return
        alpha_elem = srgb.find(qn("a:alpha"))
        if alpha_elem is None:
            alpha_elem = etree.SubElement(srgb, qn("a:alpha"))
        alpha_elem.set("val", str(alpha))
    except Exception as exc:  # noqa: BLE001
        logger.warning("fill opacity apply failed: %s", exc)


def _set_bullet(paragraph, bullet_char: str | None) -> None:
    try:
        p_pr = paragraph._p.get_or_add_pPr()
        for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
            existing = p_pr.find(qn(tag))
            if existing is not None:
                p_pr.remove(existing)

        bullet_font = p_pr.find(qn("a:buFont"))
        if bullet_font is None:
            bullet_font = etree.SubElement(p_pr, qn("a:buFont"))
        bullet_font.set("typeface", "Arial")

        bullet_char_el = etree.SubElement(p_pr, qn("a:buChar"))
        bullet_char_el.set("char", (bullet_char or "•")[0])
    except Exception as exc:  # noqa: BLE001
        logger.warning("bullet apply failed: %s", exc)


def _set_line_dash(connector, line_style: LineStyle) -> None:
    try:
        dash_map = {
            LineStyle.DASHED: "dash",
            LineStyle.DOTTED: "dot",
        }
        if line_style not in dash_map:
            return
        line = connector.line._get_or_add_ln()
        existing = line.find(qn("a:prstDash"))
        if existing is not None:
            line.remove(existing)
        dash = etree.SubElement(line, qn("a:prstDash"))
        dash.set("val", dash_map[line_style])
    except Exception as exc:  # noqa: BLE001
        logger.warning("line dash apply failed: %s", exc)


def _set_arrow(connector, arrow_type: ArrowType, *, tail: bool) -> None:
    arrow_map = {
        ArrowType.TRIANGLE: ("triangle", "med", "med"),
        ArrowType.STEALTH: ("stealth", "med", "med"),
        ArrowType.OPEN: ("arrow", "med", "med"),
    }
    if arrow_type not in arrow_map:
        return
    try:
        arrow_kind, width, length = arrow_map[arrow_type]
        line = connector.line._get_or_add_ln()
        tag = qn("a:tailEnd") if tail else qn("a:headEnd")
        existing = line.find(tag)
        if existing is not None:
            line.remove(existing)
        node = etree.SubElement(line, tag)
        node.set("type", arrow_kind)
        node.set("w", width)
        node.set("len", length)
    except Exception as exc:  # noqa: BLE001
        logger.warning("arrow apply failed: %s", exc)


def _apply_paragraphs(text_frame, paragraphs: list[TextParagraph]) -> None:
    text_frame.word_wrap = True
    if not paragraphs:
        return

    for index, paragraph_data in enumerate(paragraphs):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        if index == 0:
            for run in list(paragraph.runs):
                run._r.getparent().remove(run._r)

        paragraph.alignment = ALIGN_MAP.get(paragraph_data.alignment, PP_ALIGN.LEFT)
        if paragraph_data.bullet:
            paragraph.level = paragraph_data.indent_level
            _set_bullet(paragraph, paragraph_data.bullet_char)
        if paragraph_data.line_spacing and paragraph_data.line_spacing != 1.0:
            try:
                paragraph.line_spacing = paragraph_data.line_spacing
            except Exception:
                pass

        for run_data in paragraph_data.runs:
            run = paragraph.add_run()
            run.text = run_data.text or ""
            _apply_run_font(run, run_data)
            if run_data.hyperlink:
                try:
                    run.hyperlink.address = run_data.hyperlink
                except Exception as exc:  # noqa: BLE001
                    logger.warning("hyperlink apply failed: %s", exc)


def _apply_run_font(run, run_data) -> None:
    font = run.font
    try:
        font.name = resolve_font(run_data.font.family)
        font.size = Pt(run_data.font.size_pt)
        font.bold = run_data.font.bold
        font.italic = run_data.font.italic
        font.underline = run_data.font.underline
        font.color.rgb = hex_to_rgb(run_data.font.color_hex)
    except Exception as exc:  # noqa: BLE001
        logger.warning("font apply failed: %s", exc)


def _build_text_box(slide, elem: SlideElement, w_in: float, h_in: float) -> None:
    left, top, width, height = _bbox_to_emu(elem, w_in, h_in)
    text_box = slide.shapes.add_textbox(left, top, width, height)
    text_box.rotation = elem.rotation
    _apply_paragraphs(text_box.text_frame, elem.paragraphs)


def _build_shape(slide, elem: SlideElement, w_in: float, h_in: float) -> None:
    shape_kind = elem.shape_kind or ShapeKind.RECTANGLE
    mso_shape = SHAPE_KIND_TO_MSO.get(shape_kind, MSO_SHAPE.RECTANGLE)
    left, top, width, height = _bbox_to_emu(elem, w_in, h_in)
    shape = slide.shapes.add_shape(mso_shape, left, top, width, height)
    shape.rotation = elem.rotation

    if elem.fill_color_hex:
        shape.fill.solid()
        shape.fill.fore_color.rgb = hex_to_rgb(elem.fill_color_hex)
        if elem.fill_opacity < 1.0:
            _set_fill_opacity(shape, elem.fill_opacity)
    else:
        shape.fill.background()

    if elem.border_color_hex:
        shape.line.color.rgb = hex_to_rgb(elem.border_color_hex)
        shape.line.width = Pt(elem.border_width_pt or 1.0)
    else:
        try:
            shape.line.fill.background()
        except Exception:
            pass

    if elem.paragraphs:
        _apply_paragraphs(shape.text_frame, elem.paragraphs)


def _build_line(slide, elem: SlideElement, w_in: float, h_in: float) -> None:
    x1 = pct_to_emu(elem.bbox.x_pct, w_in)
    y1 = pct_to_emu(elem.bbox.y_pct, h_in)
    x2 = pct_to_emu(elem.bbox.x_pct + elem.bbox.width_pct, w_in)
    y2 = pct_to_emu(elem.bbox.y_pct + elem.bbox.height_pct, h_in)
    connector = slide.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT, x1, y1, x2, y2)
    connector.line.color.rgb = hex_to_rgb(elem.line_color_hex or "#000000")
    connector.line.width = Pt(elem.line_width_pt)
    if elem.line_style != LineStyle.SOLID:
        _set_line_dash(connector, elem.line_style)
    if elem.end_arrow != ArrowType.NONE:
        _set_arrow(connector, elem.end_arrow, tail=True)
    if elem.start_arrow != ArrowType.NONE:
        _set_arrow(connector, elem.start_arrow, tail=False)


def _build_image_placeholder(slide, elem: SlideElement, w_in: float, h_in: float) -> None:
    left, top, width, height = _bbox_to_emu(elem, w_in, h_in)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_to_rgb(elem.image_placeholder_color or "#E0E0E0")
    shape.line.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    shape.line.width = Pt(0.5)
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = f"[Image: {elem.image_description or 'unknown'}]"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    run.font.italic = True


def _build_table(slide, elem: SlideElement, w_in: float, h_in: float) -> None:
    if not elem.table_rows or not elem.table_cols or not elem.table_data:
        return
    left, top, width, height = _bbox_to_emu(elem, w_in, h_in)
    table_shape = slide.shapes.add_table(elem.table_rows, elem.table_cols, left, top, width, height)
    table = table_shape.table
    for row_index, row_data in enumerate(elem.table_data):
        if row_index >= elem.table_rows:
            break
        for col_index, cell_text in enumerate(row_data):
            if col_index >= elem.table_cols:
                break
            table.cell(row_index, col_index).text = str(cell_text)


def _build_raster_asset(
    slide,
    elem: SlideElement,
    w_in: float,
    h_in: float,
    asset_lookup: dict[str, str] | None,
) -> bool:
    if not elem.asset_id or not asset_lookup:
        return False
    asset_path = asset_lookup.get(elem.asset_id)
    if not asset_path or not Path(asset_path).exists():
        return False
    left, top, width, height = _bbox_to_emu(elem, w_in, h_in)
    slide.shapes.add_picture(asset_path, left, top, width, height)
    return True


def _set_slide_background(slide, hex_color: str) -> None:
    try:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = hex_to_rgb(hex_color)
    except Exception as exc:  # noqa: BLE001
        logger.warning("background apply failed: %s", exc)


BUILDER_MAP = {
    ElementType.TEXT_BOX: _build_text_box,
    ElementType.SHAPE: _build_shape,
    ElementType.LINE: _build_line,
    ElementType.CONNECTOR: _build_line,
    ElementType.IMAGE: _build_image_placeholder,
    ElementType.TABLE: _build_table,
}


def build_presentation(
    data: PresentationData,
    *,
    asset_lookup: dict[str, str] | None = None,
) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(data.slide_width_inches)
    prs.slide_height = Inches(data.slide_height_inches)
    blank_layout = prs.slide_layouts[6]

    for slide_data in data.slides:
        slide = prs.slides.add_slide(blank_layout)
        if slide_data.background_color_hex and slide_data.background_color_hex != "#FFFFFF":
            _set_slide_background(slide, slide_data.background_color_hex)

        sorted_elements = sorted(slide_data.elements, key=lambda element: element.z_order)
        for element in sorted_elements:
            try:
                if element.render_strategy != RenderStrategy.PRIMITIVE and _build_raster_asset(
                    slide,
                    element,
                    data.slide_width_inches,
                    data.slide_height_inches,
                    asset_lookup,
                ):
                    continue

                builder = BUILDER_MAP.get(element.element_type)
                if builder is None:
                    logger.warning(
                        "Unsupported element type %s on slide %s",
                        element.element_type,
                        slide_data.slide_number,
                    )
                    continue
                builder(slide, element, data.slide_width_inches, data.slide_height_inches)
            except Exception as exc:  # noqa: BLE001
                logger.warning(
                    "Slide %s element %s (%s) failed: %s",
                    slide_data.slide_number,
                    element.id,
                    element.element_type,
                    exc,
                )

        if slide_data.notes:
            try:
                slide.notes_slide.notes_text_frame.text = slide_data.notes
            except Exception as exc:  # noqa: BLE001
                logger.warning("notes apply failed: %s", exc)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def build_pptx(
    data: PresentationData,
    *,
    asset_lookup: dict[str, str] | None = None,
) -> bytes:
    return build_presentation(data, asset_lookup=asset_lookup)
