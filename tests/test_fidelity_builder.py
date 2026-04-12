from __future__ import annotations

import io

import pytest
from PIL import Image
from pptx import Presentation

from app.config import settings
from app.fidelity_assets import ensure_assets_for_slide
from app.models import BoundingBox, ElementType, FinalElement, PresentationData, RenderStrategy, SlideData
from app.pptx_builder import build_presentation
from app.session_store import create_session, get_asset_path, save_final_elements, save_source_image


def _make_png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (400, 225), color="#89A7FF").save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.fixture()
def tmp_output_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "output_dir", tmp_path)
    yield tmp_path


def test_ensure_assets_for_slide_creates_crop_asset(tmp_output_dir):
    session = create_session(total_slides=1)
    save_source_image(session.job_id, 1, _make_png())
    final = FinalElement(
        final_id="final-raster",
        element_type=ElementType.IMAGE,
        bbox=BoundingBox(x_pct=10, y_pct=10, width_pct=30, height_pct=25),
        source_region=BoundingBox(x_pct=10, y_pct=10, width_pct=30, height_pct=25),
        render_strategy=RenderStrategy.RASTER,
        image_description="diagram",
    )

    updated = ensure_assets_for_slide(session.job_id, 1, [final])

    assert updated[0].asset_id is not None
    assert get_asset_path(session.job_id, updated[0].asset_id).exists()


def test_build_presentation_uses_raster_asset_lookup(tmp_output_dir):
    session = create_session(total_slides=1)
    save_source_image(session.job_id, 1, _make_png())
    final = FinalElement(
        final_id="final-raster",
        element_type=ElementType.IMAGE,
        bbox=BoundingBox(x_pct=10, y_pct=10, width_pct=30, height_pct=25),
        source_region=BoundingBox(x_pct=10, y_pct=10, width_pct=30, height_pct=25),
        render_strategy=RenderStrategy.RASTER,
        image_description="diagram",
    )
    updated = ensure_assets_for_slide(session.job_id, 1, [final])
    save_final_elements(session.job_id, 1, updated)

    asset_path = get_asset_path(session.job_id, updated[0].asset_id)
    presentation = PresentationData(
        slides=[
            SlideData(
                slide_number=1,
                elements=[updated[0]],
            )
        ]
    )

    pptx_bytes = build_presentation(
        presentation,
        asset_lookup={updated[0].asset_id: str(asset_path)},
    )
    prs = Presentation(io.BytesIO(pptx_bytes))

    assert len(prs.slides) == 1
    assert len(prs.slides[0].shapes) >= 1
